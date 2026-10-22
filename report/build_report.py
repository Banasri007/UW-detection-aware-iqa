"""Build the FCV Lab mini-project report from the department template.

    python report/build_report.py            # -> report/FCV_MiniProject_Report.docx (+ .pdf via Word)

Starts from report/template/FCV_MiniProject_Template.docx (the faculty template converted
to .docx), fills the cover page and certificate, replaces the static contents list with a
real Word table of contents, and appends the seven template chapters. Code listings in
Chapter 5 are pulled from the project source with inspect.getsource, so they always match
the code that produced the results. Figures come from report/assets/.
"""
from __future__ import annotations

import inspect
import re
import sys
import textwrap
from pathlib import Path

import numpy as np
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
ASSETS = HERE / "assets"
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "app"))

FONT = "Garamond"
TITLE = ("Seeing Better Is Not Detecting Better: Benchmarking No-Reference Quality Metrics and "
         "Learning Detection-Aware Enhancement Decisions for Underwater Images")
TEAM = [("Akshant Beuria", "240962474", "36"),
        ("Javvaji Aditya Banasri", "240962478", "37"),
        ("Mansi Sinha", "240962484", "39")]
SECTION = "____"  # fill in your section

# ===================================================================== references (IEEE)
REFS = {
    "gonzalez2012": 'R. C. Gonzalez and R. E. Woods, *Digital Image Processing*, 3rd ed. Pearson Education, 2012.',
    "szeliski2011": 'R. Szeliski, *Computer Vision: Algorithms and Applications*. Springer, 2011.',
    "panetta2016uiqm": 'K. Panetta, C. Gao and S. Agaian, "Human-visual-system-inspired underwater image quality measures," *IEEE J. Oceanic Eng.*, vol. 41, no. 3, pp. 541–551, 2016.',
    "yang2015uciqe": 'M. Yang and A. Sowmya, "An underwater color image quality evaluation metric," *IEEE Trans. Image Process.*, vol. 24, no. 12, pp. 6062–6071, 2015.',
    "li2020uieb": 'C. Li, C. Guo, W. Ren, R. Cong, J. Hou, S. Kwong and D. Tao, "An underwater image enhancement benchmark dataset and beyond," *IEEE Trans. Image Process.*, vol. 29, pp. 4376–4389, 2020.',
    "li2022limits": 'C. Y. Li and A. Cavallaro, "On the limits of perceptual quality measures for enhanced underwater images," in *Proc. IEEE ICIP*, 2022, pp. 4148–4152.',
    "palazari2026reliable": 'S. Palazari and E. Dumic, "Towards reliable evaluation of underwater image enhancement using subjective and objective analysis," *Electronics*, vol. 15, no. 11, art. 2412, 2026.',
    "wang2024isuie": 'Y. Wang, J. Guo, W. He, H. Gao, H. Yue, Z. Zhang and C. Li, "Is underwater image enhancement all object detectors need?," *IEEE J. Oceanic Eng.*, vol. 49, no. 2, pp. 606–621, 2024.',
    "saleem2025influence": 'A. Saleem, A. Awad, S. Paheding, E. Lucas, T. C. Havens and P. C. Esselman, "Understanding the influence of image enhancement on underwater object detection: a quantitative and qualitative study," *Remote Sensing*, vol. 17, no. 2, art. 185, 2025.',
    "awad2026revisiting": 'A. Awad, A. Saleem, S. Paheding, E. Lucas, S. Al-Ratrout and T. C. Havens, "Revisiting underwater image enhancement for object detection: a unified quality–detection evaluation framework," *J. Imaging*, vol. 12, no. 1, art. 18, 2026.',
    "hou2023uid2021": 'G. Hou, Y. Li, H. Yang, K. Li and Z. Pan, "UID2021: an underwater image dataset for evaluation of no-reference quality assessment metrics," *ACM Trans. Multimedia Comput. Commun. Appl.*, vol. 19, no. 4, pp. 1–24, 2023.',
    "jiang2022saud": 'Q. Jiang, Y. Gu, C. Li, R. Cong and F. Shao, "Underwater image enhancement quality evaluation: benchmark dataset and objective metric," *IEEE Trans. Circuits Syst. Video Technol.*, vol. 32, no. 9, pp. 5959–5974, 2022.',
    "guo2023uranker": 'C. Guo, R. Wu, X. Jin, L. Han, W. Zhang, Z. Chai and C. Li, "Underwater Ranker: learn which is better and how to be better," in *Proc. AAAI*, vol. 37, no. 1, 2023, pp. 702–709.',
    "fu2023ruod": 'C. Fu, R. Liu, X. Fan, P. Chen, H. Fu, W. Yuan, M. Zhu and Z. Luo, "Rethinking general underwater object detection: datasets, challenges, and solutions," *Neurocomputing*, vol. 517, pp. 243–256, 2023.',
    "liu2021duo": 'C. Liu, H. Li, S. Wang, M. Zhu, D. Wang, X. Fan and Z. Wang, "A dataset and benchmark of underwater object detection for robot picking," in *Proc. IEEE ICMEW*, 2021, pp. 1–6.',
    "odverse33": 'T. Jiang and Y. Zhong, "ODverse33: is the new YOLO version always better? A multi domain benchmark from YOLO v5 to v11," arXiv:2502.14314, 2025.',
    "zuiderveld1994clahe": 'K. Zuiderveld, "Contrast limited adaptive histogram equalization," in *Graphics Gems IV*. Academic Press, 1994, pp. 474–485.',
    "buchsbaum1980grayworld": 'G. Buchsbaum, "A spatial processor model for object colour perception," *J. Franklin Inst.*, vol. 310, no. 1, pp. 1–26, 1980.',
    "drews2013udcp": 'P. Drews Jr., E. do Nascimento, F. Moraes, S. Botelho and M. Campos, "Transmission estimation in underwater single images," in *Proc. IEEE ICCV Workshops*, 2013, pp. 825–830.',
    "ancuti2018fusion": 'C. O. Ancuti, C. Ancuti, C. De Vleeschouwer and P. Bekaert, "Color balance and fusion for underwater image enhancement," *IEEE Trans. Image Process.*, vol. 27, no. 1, pp. 379–393, 2018.',
    "islam2020funie": 'M. J. Islam, Y. Xia and J. Sattar, "Fast underwater image enhancement for improved visual perception," *IEEE Robot. Autom. Lett.*, vol. 5, no. 2, pp. 3227–3234, 2020.',
    "mittal2013niqe": 'A. Mittal, R. Soundararajan and A. C. Bovik, "Making a \'completely blind\' image quality analyzer," *IEEE Signal Process. Lett.*, vol. 20, no. 3, pp. 209–212, 2013.',
    "mittal2012brisque": 'A. Mittal, A. K. Moorthy and A. C. Bovik, "No-reference image quality assessment in the spatial domain," *IEEE Trans. Image Process.*, vol. 21, no. 12, pp. 4695–4708, 2012.',
    "ke2021musiq": 'J. Ke, Q. Wang, Y. Wang, P. Milanfar and F. Yang, "MUSIQ: multi-scale image quality transformer," in *Proc. IEEE/CVF ICCV*, 2021, pp. 5148–5157.',
    "wang2023clipiqa": 'J. Wang, K. C. K. Chan and C. C. Loy, "Exploring CLIP for assessing the look and feel of images," in *Proc. AAAI*, vol. 37, no. 2, 2023, pp. 2555–2563.',
    "chen2024topiq": 'C. Chen, J. Mo, J. Hou, H. Wu, L. Liao, W. Sun, Q. Yan and W. Lin, "TOPIQ: a top-down approach from semantics to distortions for image quality assessment," *IEEE Trans. Image Process.*, vol. 33, pp. 2404–2418, 2024.',
    "zhang2023liqe": 'W. Zhang, G. Zhai, Y. Wei, X. Yang and K. Ma, "Blind image quality assessment via vision-language correspondence: a multitask learning perspective," in *Proc. IEEE/CVF CVPR*, 2023, pp. 14071–14081.',
    "yang2022maniqa": 'S. Yang et al., "MANIQA: multi-dimension attention network for no-reference image quality assessment," in *Proc. IEEE/CVF CVPR Workshops*, 2022, pp. 1190–1199.',
    "golestaneh2022tres": 'S. A. Golestaneh, S. Dadsetan and K. M. Kitani, "No-reference image quality assessment via transformers, relative ranking, and self-consistency," in *Proc. IEEE/CVF WACV*, 2022, pp. 3989–3999.',
    "agnolucci2024arniqa": 'L. Agnolucci, L. Galteri, M. Bertini and A. Del Bimbo, "ARNIQA: learning distortion manifold for image quality assessment," in *Proc. IEEE/CVF WACV*, 2024, pp. 188–197.',
    "agnolucci2024qualiclip": 'L. Agnolucci, L. Galteri and M. Bertini, "Quality-aware image-text alignment for opinion-unaware image quality assessment," arXiv:2403.11176, 2024.',
    "chen2022iqapytorch": 'C. Chen and J. Mo, "IQA-PyTorch: PyTorch toolbox for image quality assessment," 2022. [Online]. Available: https://github.com/chaofengc/IQA-PyTorch',
    "jocher2023ultralytics": 'G. Jocher and J. Qiu, "Ultralytics YOLO," 2024. [Online]. Available: https://github.com/ultralytics/ultralytics',
    "lin2014coco": 'T.-Y. Lin et al., "Microsoft COCO: common objects in context," in *Proc. ECCV*, 2014, pp. 740–755.',
    "radford2021clip": 'A. Radford et al., "Learning transferable visual models from natural language supervision," in *Proc. ICML*, 2021.',
    "oquab2024dinov2": 'M. Oquab et al., "DINOv2: learning robust visual features without supervision," *Trans. Mach. Learn. Res.*, 2024.',
    "he2016resnet": 'K. He, X. Zhang, S. Ren and J. Sun, "Deep residual learning for image recognition," in *Proc. IEEE CVPR*, 2016, pp. 770–778.',
    "vqeg2000": 'Video Quality Experts Group, "Final report from the Video Quality Experts Group on the validation of objective models of video quality assessment," VQEG, 2000.',
    "efron1994bootstrap": 'B. Efron and R. J. Tibshirani, *An Introduction to the Bootstrap*. Chapman & Hall/CRC, 1994.',
    "canny1986": 'J. Canny, "A computational approach to edge detection," *IEEE Trans. Pattern Anal. Mach. Intell.*, vol. PAMI-8, no. 6, pp. 679–698, 1986.',
    "otsu1979": 'N. Otsu, "A threshold selection method from gray-level histograms," *IEEE Trans. Syst., Man, Cybern.*, vol. 9, no. 1, pp. 62–66, 1979.',
    "lowe2004sift": 'D. G. Lowe, "Distinctive image features from scale-invariant keypoints," *Int. J. Comput. Vis.*, vol. 60, no. 2, pp. 91–110, 2004.',
    "rosten2006fast": 'E. Rosten and T. Drummond, "Machine learning for high-speed corner detection," in *Proc. ECCV*, 2006, pp. 430–443.',
    "harris1988": 'C. Harris and M. Stephens, "A combined corner and edge detector," in *Proc. Alvey Vision Conference*, 1988, pp. 147–151.',
    "dalal2005hog": 'N. Dalal and B. Triggs, "Histograms of oriented gradients for human detection," in *Proc. IEEE CVPR*, 2005, pp. 886–893.',
    "ojala2002lbp": 'T. Ojala, M. Pietikäinen and T. Mäenpää, "Multiresolution gray-scale and rotation invariant texture classification with local binary patterns," *IEEE Trans. Pattern Anal. Mach. Intell.*, vol. 24, no. 7, pp. 971–987, 2002.',
    "fastapi": 'S. Ramírez, "FastAPI," 2018. [Online]. Available: https://fastapi.tiangolo.com',
    "opencv": 'G. Bradski, "The OpenCV library," *Dr. Dobb\'s Journal of Software Tools*, 2000.',
}


def src(obj) -> str:
    return textwrap.dedent(inspect.getsource(obj)).strip()


def code_lines(obj, start: str | None = None, end: str | None = None) -> str:
    """Excerpt of an object's source between two marker substrings (inclusive)."""
    lines = src(obj).splitlines()
    i = next((k for k, l in enumerate(lines) if start in l), 0) if start else 0
    j = next((k for k, l in enumerate(lines) if k >= i and end in l), len(lines) - 1) if end else len(lines) - 1
    return "\n".join(lines[i:j + 1])


# ===================================================================== content
def content():
    from uwiqa.data import dedup
    from uwiqa.detect import metrics as dm
    from uwiqa.enhance import classical, registry, sweeps
    import importlib
    uiqm_mod = importlib.import_module("uwiqa.metrics.uiqm")  # the package re-exports a function named uiqm
    import vision as V

    B = []
    H1 = lambda t: B.append(("h1", t))
    H2 = lambda t: B.append(("h2", t))
    P = lambda t: B.append(("p", t))
    BUL = lambda items: B.append(("bullets", items))
    TAB = lambda cap, head, rows, widths=None, size=10: B.append(("table", cap, head, rows, widths, size))
    FIG = lambda path, cap, w=6.3: B.append(("fig", path, cap, w))
    CODE = lambda cap, code: B.append(("code", cap, code))

    # ----------------------------------------------------------------- abstract
    B.append(("h1u", "ABSTRACT"))
    P("Enhanced underwater images are routinely judged by the hand-crafted no-reference (NR) quality metrics UIQM and "
      "UCIQE, and enhancement is routinely applied before object detection. This project tests both habits on public "
      "data under one reproducible protocol. First, 16 NR image-quality metrics are benchmarked against human opinion "
      "scores on UID2021 (960 images, 77 observers) with bootstrap confidence intervals, paired significance tests and "
      "intra-scene rank correlation. Only the underwater-trained URanker reaches a Spearman correlation of 0.62; in-air "
      "deep metrics fall to near zero when ranking enhancements of the same scene, and UCIQE and URanker are inflated by "
      "deliberate over-enhancement. Second, every held-out image of a leakage-safe re-split of RUOD, and of a "
      "duplicate-free subset of DUO, is labelled with the change in per-image detection AP caused by five enhancers, "
      "measured by a frozen YOLO11-s detector and calibrated against a label-noise floor. All five enhancers lower mean "
      "AP (by 0.02–0.08 on RUOD and 0.05–0.20 on DUO), yet a per-image oracle could gain 0.025 mAP, about ten times what "
      "label noise alone produces. A lightweight predictor that sees only the raw image ranks the detection-utility "
      "change far better than any NR-IQA metric (Spearman 0.32 vs. at most 0.13), but selective enhancement driven by it "
      "at best matches never enhancing, whereas choosing the enhancer with the best NR-IQA gain, as common practice "
      "suggests, costs up to 0.17 mAP. The project also shows that the public re-split of RUOD leaks near-duplicate "
      "frames across its splits and overlaps heavily with DUO. An interactive web application demonstrates the "
      "enhancers, quality scores, Lab 1–5 image operations and detections on individual images.")
    P("**Keywords:** underwater image enhancement, no-reference image quality assessment, object detection, "
      "task-aware quality, benchmark, OpenCV.")

    # ----------------------------------------------------------------- chapter 1
    H1("CHAPTER 1: INTRODUCTION")
    H2("1.1 Background")
    P("Light travelling through water is absorbed in a wavelength-dependent way: long (red) wavelengths are absorbed "
      "first, then orange and yellow, which gives underwater photographs their characteristic blue-green cast. Suspended "
      "particles also scatter light. Forward scattering blurs object edges, and backscatter adds a veiling haze that "
      "suppresses contrast. The result is low contrast, strong colour casts, reduced visibility and non-uniform "
      "illumination, all of which worsen with depth, turbidity and camera-to-object distance.")
    P("Underwater image enhancement (UIE) tries to undo these effects. It matters for marine ecology surveys, coral reef "
      "monitoring, aquaculture, autonomous and remotely operated underwater vehicles, archaeology and debris "
      "monitoring. In almost all of these applications the enhanced image is not the final product; it is the input to "
      "an automated task such as object detection.")
    H2("1.2 Underwater image enhancement")
    P("UIE methods fall into three broad families. *Classical* methods manipulate intensities and colours directly, for "
      "example contrast-limited adaptive histogram equalisation (CLAHE) [@zuiderveld1994clahe] and gray-world white "
      "balance [@buchsbaum1980grayworld], the same operations studied in Lab 2. *Physics-based* methods invert an "
      "image-formation model, such as the underwater dark channel prior (UDCP) [@drews2013udcp]. *Fusion and "
      "learning-based* methods combine several enhanced inputs [@ancuti2018fusion] or learn the mapping from data, for "
      "example the FUnIE-GAN generator [@islam2020funie].")
    H2("1.3 Evaluating enhanced images")
    P("Full-reference metrics such as PSNR and SSIM compare an image with a clean reference, but no clean version of a "
      "real underwater scene exists: benchmark references such as those of UIEB were chosen by human voting among "
      "candidate enhancements [@li2020uieb]. The field therefore relies on *no-reference* (NR) metrics, above all UIQM "
      "[@panetta2016uiqm] and UCIQE [@yang2015uciqe]. Both are weighted sums of low-level colourfulness, sharpness, "
      "contrast and saturation statistics, and both are known to reward over-processed results [@li2020uieb, "
      "@li2022limits]. Modern deep NR models trained on in-air distortions are a possible replacement, but the evidence "
      "for their transfer to underwater images is thin [@palazari2026reliable].")
    H2("1.4 Enhancement for downstream detection")
    P("Studies that train detectors on raw and enhanced data find that enhancement does not reliably improve detection "
      "[@wang2024isuie], that over-enhancement in particular degrades it [@saleem2025influence], and that an oracle "
      "choosing per image between raw and enhanced data would beat both [@awad2026revisiting]. What is missing is a "
      "metric that predicts, for a given image, whether enhancement will help the detector.")
    H2("1.5 Contributions")
    BUL(["A benchmark of 16 NR-IQA metrics (hand-crafted, in-air deep and underwater-trained) against human opinion "
         "scores on UID2021, with bootstrap confidence intervals, paired significance tests, intra-scene correlation and "
         "an over-enhancement stress test that shows which metrics can be gamed.",
         "A near-duplicate analysis of the RUOD and DUO detection datasets that exposes leakage across the public "
         "splits and between the datasets, and leakage-safe splits built from it.",
         "Per-image detection-utility labels for five enhancers on 7,000 held-out RUOD images and 4,539 duplicate-free "
         "DUO images, measured with a frozen detector and calibrated against a label-noise floor.",
         "A detection-aware predictor that sees only the raw image, compared with NR-IQA metrics as utility predictors, "
         "and a dataset-level evaluation of selective-enhancement policies with an oracle upper bound and a noise-only "
         "control.",
         "An interactive web demo application that applies the enhancers, quality metrics, Lab 1–5 operations and "
         "the trained detector to individual images."])
    H2("1.6 Organisation of the report")
    P("Chapter 2 states the problem, objectives and system requirements. Chapter 3 describes the methodology and the "
      "demo application. Chapter 4 presents the results with snapshots of the experiments and the application. "
      "Chapter 5 maps the functions from Labs 1–5 to where they are used in the project. Chapter 6 concludes and "
      "discusses limitations and future work, and Chapter 7 lists the references.")

    # ----------------------------------------------------------------- chapter 2
    H1("CHAPTER 2: PROBLEM STATEMENT & OBJECTIVES")
    H2("2.1 Problem statement")
    P("The de facto standard NR quality metrics for enhanced underwater images, UIQM and UCIQE, are hand-crafted linear "
      "combinations of low-level colourfulness, sharpness, contrast and saturation statistics. They correlate weakly "
      "with human perceptual judgement and can be inflated, or effectively gamed, by over-enhancement. Furthermore, no "
      "existing NR metric predicts whether enhancing a given underwater image will actually improve a downstream task "
      "such as object detection, even though downstream utility is what most operational underwater vision pipelines "
      "ultimately care about.")
    P("The project therefore addresses two coupled questions: (1) how well do classical and modern learning-based NR-IQA "
      "models track human judgement of enhanced underwater images, measured under one consistent protocol; and (2) can "
      "a no-reference model predict, per image and from the raw image alone, whether enhancement will help or hurt "
      "object detection?")
    H2("2.2 Research gaps")
    BUL(["**Gap 1: UIQM and UCIQE are unreliable** and reward over-enhancement [@li2020uieb, @li2022limits].",
         "**Gap 2: transfer of in-air deep NR-IQA to underwater images is unproven at scale**; the strongest positive "
         "evidence comes from two small datasets and four enhancers [@palazari2026reliable].",
         "**Gap 3: no metric predicts downstream task utility.** Prior work analyses detection gains after the fact with "
         "an oracle [@awad2026revisiting] but offers no deployable predictor.",
         "**Gap 4: underwater subjective datasets are fragmented**, with different scales and protocols, so published "
         "correlations are hard to compare."])
    H2("2.3 Objectives")
    TAB("Objectives and how each was measured.", ["ID", "Objective", "Measured by", "Outcome"], [
        ["O1", "NR-IQA benchmark against human scores", "SRCC / PLCC / KRCC with bootstrap CIs, intra-scene SRCC",
         "Done on UID2021 for 16 metrics"],
        ["O2", "Over-enhancement stress test", "Score vs. distortion strength (Spearman)",
         "Done: 4 distortions × 7 strengths × 60 images"],
        ["O3", "Per-image detection-utility labels", "ΔAP per image and enhancer, frozen detector",
         "Done: 7,000 RUOD + 4,539 DUO images × 5 enhancers"],
        ["O4", "Detection-aware NR predictor", "SRCC vs. true ΔAP; AUROC, balanced accuracy, F1",
         "Done: 4 backbones × 3 heads vs. 10 NR-IQA baselines"],
        ["O5", "Cross-dataset generalisation", "Same metrics on DUO, nothing tuned on DUO",
         "Done on duplicate-free DUO"],
        ["O6", "Selective enhancement (practical validation)", "Dataset mAP of per-image policies",
         "Done, with oracle and noise controls"],
        ["O7", "Interactive demo application", "Working UI on real samples", "Done (web app)"],
    ], widths=[0.45, 2.0, 2.2, 1.85], size=9.5)
    H2("2.4 Changes from the synopsis")
    P("Some plans from the Phase-I synopsis were adapted during the project, for reasons of data access, software "
      "compatibility and time:")
    BUL(["**Subjective datasets:** the benchmark uses UID2021 only, instead of three datasets. UID2021 offers the most "
         "complete public MOS (15 enhancers per scene); Gap 4 is therefore only partly addressed.",
         "**Enhancers:** Water-Net and Ucolor are available only for TensorFlow 1.x, which does not run on current "
         "Kaggle images; CLAHE and gray-world were used instead, alongside UDCP, fusion and FUnIE-GAN.",
         "**Detection data:** RUOD and DUO were taken from the YOLO-format ODverse33 copy [@odverse33]. Brackish was "
         "dropped; DUO serves as the cross-dataset test.",
         "**Additions:** a near-duplicate analysis (which revealed split leakage), a label-noise control, and the demo "
         "application. The optional in-house subjective study (O6 in the synopsis) was not conducted."])
    H2("2.5 System requirements")
    TAB("Hardware used.", ["Component", "Specification", "Used for"], [
        ["Development machine", "Intel Core i5-13500H laptop, Windows 11, no discrete GPU",
         "Code, unit tests, demo application, report"],
        ["Kaggle GPU sessions", "2 × NVIDIA Tesla T4 (16 GB each), 4 vCPUs",
         "Detector training, utility labels, deep NR-IQA, feature extraction"],
        ["Kaggle CPU sessions", "4 vCPUs", "Data checks, duplicate search, figures, demo assets"],
        ["Storage", "≈7 GB detection data, ≈1 GB UID2021, < 20 GB working space", "Datasets and results"],
    ], widths=[1.5, 2.8, 2.2], size=9.5)
    TAB("Software used.", ["Software", "Version", "Purpose"], [
        ["Python", "3.12 (Kaggle), 3.13 (local)", "Implementation language"],
        ["OpenCV", "4.10", "Image I/O, colour spaces, enhancement, filtering, features [@opencv]"],
        ["NumPy / SciPy / pandas", "2.x / 1.15 / 2.2", "Numerics, statistics, tables"],
        ["scikit-image", "0.26", "CIELab conversion, HOG, LBP"],
        ["PyTorch / timm", "2.7–2.10 / 1.0", "FUnIE-GAN, frozen CLIP / DINOv2 / ResNet-18 backbones"],
        ["Ultralytics", "8.4", "YOLO11-s training and inference [@jocher2023ultralytics]"],
        ["IQA-PyTorch (pyiqa)", "0.1.16", "Deep NR-IQA metrics [@chen2022iqapytorch]"],
        ["scikit-learn", "1.6", "Ridge and logistic heads, AUROC, balanced accuracy"],
        ["FastAPI / Uvicorn", "0.104 / 0.x", "Demo application back end (JSON API) [@fastapi]"],
        ["HTML5, CSS3, JavaScript", "–", "Demo application front end (no framework, SVG charts)"],
        ["Matplotlib / Playwright", "3.10 / 1.x", "Figures / application screenshots"],
        ["Kaggle Notebooks, Git, GitHub, VS Code", "–", "Compute, version control, development"],
    ], widths=[2.0, 1.6, 2.9], size=9.5)
    P("The demo application runs on any recent CPU with about 8 GB of RAM; one image is fully analysed in a few "
      "seconds.")

    # ----------------------------------------------------------------- chapter 3
    H1("CHAPTER 3: METHODOLOGY")
    H2("3.1 Overview")
    P("Figure 3.1 shows the two halves of the pipeline. The *benchmark* half scores enhanced images with NR-IQA metrics "
      "and compares the scores with human opinion. The *detection* half enhances each held-out raw image in memory with "
      "five enhancers, runs a detector that was trained on raw images only, and records how much each enhancer changes "
      "the per-image AP. A predictor then learns to estimate that change from the raw image alone, and selective "
      "enhancement policies are evaluated at dataset level.")
    FIG(ASSETS / "pipeline.png", "Overall pipeline. Top: human-opinion benchmark. Middle: detection-utility labels. "
        "Bottom: deployable predictor and selective-enhancement policy.", 6.3)
    H2("3.2 Datasets")
    TAB("Datasets used.", ["Dataset", "Content", "Split / subset", "Role"], [
        ["UID2021 [@hou2023uid2021]", "60 raw images × (15 enhancers + raw) = 960 images, 512×384, MOS from 77 "
         "observers, 6 degradation types", "all 960", "O1 benchmark, O2 stress test (60 raw)"],
        ["RUOD [@fu2023ruod]", "14,000 images, 74,904 boxes, 10 classes", "re-split by duplicate group (Sec. 3.3)",
         "detector training, utility labels, predictor"],
        ["DUO [@liu2021duo]", "7,782 images, 74,515 boxes, 4 classes (same ids as RUOD 0–3)",
         "4,539 images sharing no duplicate group with RUOD", "cross-dataset test"],
    ], widths=[1.3, 2.4, 1.5, 1.3], size=9.5)
    P("RUOD and DUO were taken from the YOLO-format ODverse33 copy on Kaggle [@odverse33]. Their images are at native "
      "resolution (720×405 up to 3840×2160) and JPEG-compressed (RUOD ≈ quality 85, DUO ≈ 95).")
    H2("3.3 Near-duplicate analysis and leakage-safe splits")
    P("RUOD and DUO come from the same lab and contain frames of the same videos, so near-duplicate images can leak "
      "between training and test data. Every image was described by a 256-bit *difference hash* (dHash), a compact "
      "binary descriptor built from local pixel comparisons, in the spirit of the LBP descriptor of Lab 5:")
    BUL(["decode the image directly at reduced size and convert it to grey levels;",
         "resize it to 17×16 pixels;",
         "compare each pixel with its right-hand neighbour; the 256 comparisons (1 if brighter) form the hash;",
         "two images are near-duplicates if their hashes differ in at most 32 of 256 bits (Hamming distance);",
         "near-duplicate pairs are joined into groups by union-find (connected components)."])
    P("Visual inspection confirmed that pairs up to distance 26 are frames of the same scene; between 27 and 32 a few "
      "low-texture false matches appear, which only make groups conservatively large. RUOD was then re-split so that "
      "every duplicate group lies in exactly one split (Table 3.2).")
    TAB("Leakage-safe RUOD splits (whole duplicate groups per split).", ["Split", "Images", "Objects", "Used for"], [
        ["det-train (45%)", "6,300", "34,793", "detector training"],
        ["det-val (5%)", "700", "3,646", "detector early stopping"],
        ["pred-train (30%)", "4,200", "21,773", "predictor training"],
        ["pred-val (10%)", "1,400", "7,416", "predictor / policy tuning"],
        ["pred-test (10%)", "1,400", "7,276", "final in-domain results"],
    ], widths=[1.6, 1.0, 1.0, 2.9])
    H2("3.4 Image enhancement methods")
    BUL(["**CLAHE** [@zuiderveld1994clahe]: contrast-limited adaptive histogram equalisation of the L channel in CIELab "
         "(clip limit 2, 8×8 tiles).",
         "**Gray-world** [@buchsbaum1980grayworld]: scales each colour channel so that the three channel means become "
         "equal, removing a global colour cast.",
         "**UDCP** [@drews2013udcp]: estimates the transmission from the dark channel of the green and blue channels "
         "(a 15×15 minimum filter), refines it with a guided (box) filter and inverts the image-formation model.",
         "**Fusion** [@ancuti2018fusion]: red-channel compensation and white balance, then a gamma-corrected and a "
         "sharpened version are blended with Laplacian, saliency and saturation weights in a multi-scale "
         "(Gaussian/Laplacian pyramid) fusion.",
         "**FUnIE-GAN** [@islam2020funie]: a learned generator, run with the official code and weights at the working "
         "resolution."])
    P("The first four were implemented in this project with OpenCV; all five are applied in memory to the raw image "
      "after resizing it to a long side of 1280 pixels.")
    H2("3.5 No-reference quality metrics")
    P("**UIQM** [@panetta2016uiqm] is UIQM = 0.0282·UICM + 0.2953·UISM + 3.5753·UIConM, where UICM measures "
      "colourfulness from the red–green and yellow–blue opponent channels, UISM measures sharpness as a block-wise "
      "contrast measure (EME) of Sobel edge maps, and UIConM measures contrast as a block-wise logAMEE of the "
      "(Imax − Imin)/(Imax + Imin) contrast of 10×10 blocks. **UCIQE** [@yang2015uciqe] is "
      "0.4680·σc + 0.2745·conl + 0.2576·μs in CIELab: chroma standard deviation, luminance contrast between the brightest "
      "and darkest 1% of pixels, and mean saturation. Both were implemented from the original formulas; the UIQM block "
      "computation matches a widely used reference implementation exactly (verified by unit tests).")
    P("The 14 other metrics were taken from IQA-PyTorch [@chen2022iqapytorch] with their released weights: NIQE "
      "[@mittal2013niqe], BRISQUE [@mittal2012brisque], MUSIQ [@ke2021musiq], CLIP-IQA and CLIP-IQA+ "
      "[@wang2023clipiqa], TOPIQ-NR [@chen2024topiq], LIQE and LIQE-mix [@zhang2023liqe], MANIQA [@yang2022maniqa], "
      "TReS [@golestaneh2022tres], ARNIQA [@agnolucci2024arniqa], QualiCLIP [@agnolucci2024qualiclip] and the "
      "underwater-trained URanker [@guo2023uranker]. A second, legacy UCIQE implementation was included because "
      "published UCIQE code differs between papers.")
    H2("3.6 Benchmark protocol")
    P("For every metric, the Spearman (SRCC) and Kendall (KRCC) rank correlations with the mean opinion score (MOS) are "
      "computed, together with the Pearson correlation (PLCC) after the standard four-parameter logistic mapping "
      "[@vqeg2000]. Each coefficient gets a 95% percentile-bootstrap confidence interval [@efron1994bootstrap], and each "
      "metric is compared with the best one by a paired bootstrap test. Because UID2021 contains 16 versions of every "
      "scene, the mean **intra-scene SRCC** is also reported: the correlation within each scene's 16 images, which "
      "measures what a metric is actually used for, ranking enhancements of the same image.")
    H2("3.7 Over-enhancement stress test")
    P("Four controlled distortions are applied to the 60 raw UID2021 images at strengths s ∈ {0, 0.25, 0.5, 1, 1.5, 2, "
      "3}: a linear **contrast stretch** around the mean by a factor (1 + s); a **saturation boost** of the HSV "
      "saturation by (1 + s); **unsharp masking** with weight s; and **red over-compensation** (R ← R·(1 + s) + 40·s). "
      "For each metric, the Spearman correlation between its score and s is averaged over the images. A metric whose "
      "score keeps rising as the image is visibly over-processed can be gamed.")
    H2("3.8 Reference detector")
    TAB("Reference detector configuration.", ["Setting", "Value"], [
        ["Model", "YOLO11-s (Ultralytics), COCO-pretrained initialisation"],
        ["Training data", "det-train (6,300 raw RUOD images), early stopping on det-val"],
        ["Image size / epochs / batch", "640 px / 100 epochs (patience 20) / 32"],
        ["Hardware / time", "2 × Tesla T4 (DDP) / 1.97 h"],
        ["Protocol", "trained on raw images only, then frozen (isolates the effect of preprocessing)"],
    ], widths=[2.0, 4.5])
    H2("3.9 Detection-utility labels")
    P("For every held-out image x and enhancer m, the raw and enhanced images are both passed to the frozen detector, "
      "and each is scored against the ground-truth boxes with the COCO-style per-image average precision AP50:95 "
      "[@lin2014coco]. The utility label is ΔAPm(x) = AP(m(x)) − AP(x). All detections are stored, so the dataset-level "
      "mAP of any per-image choice can be recomputed without running the detector again; this AP implementation "
      "matches Ultralytics' mAP to within 0.005 on identical predictions.")
    P("Per-image AP is noisy, so a **null enhancer** was added: a visually invisible JPEG re-encode at quality 95. The "
      "noise threshold τ is the 95th percentile of |ΔAPnull| on pred-train (τ = 0.050), and an enhancer is said to "
      "*help* an image only if ΔAP > τ.")
    H2("3.10 Detection-aware predictor")
    P("The predictor sees only the raw image, resized to 224×224 and passed through a frozen backbone: CLIP ViT-B/32 "
      "[@radford2021clip], DINOv2 ViT-S/14 [@oquab2024dinov2], ResNet-18 [@he2016resnet] or CLIP+DINOv2. Three heads "
      "predict, for every enhancer, the change ΔAP and the probability that it helps: ridge regression, per-enhancer "
      "logistic regression, and a small multi-task neural network (one hidden layer of 256 units) trained with a "
      "Smooth-L1 loss on ΔAP, a class-balanced cross-entropy on 'helps' and a pairwise ranking loss, as a 3-seed "
      "ensemble with early stopping. It is trained on pred-train, tuned on pred-val and evaluated on pred-test and DUO.")
    P("As baselines, NR-IQA metrics are used directly as utility predictors, either through the quality of the raw image "
      "(−Q(x): 'enhance images that look bad') or through the quality gain of the enhancement (Q(m(x)) − Q(x)), for "
      "UIQM, UCIQE, TOPIQ, LIQE and URanker.")
    H2("3.11 Selective-enhancement policies")
    P("A policy picks, for each image, the enhancer m* with the highest predicted score and applies it only if that score "
      "exceeds a margin δ; otherwise it keeps the raw image. The margin and the model are chosen on pred-val only, "
      "either for mean per-image AP or for dataset mAP. Policies are compared by dataset mAP50:95 with always-raw, "
      "always-m, the common practice of applying the enhancer with the largest NR-IQA gain, a per-image *oracle* (best "
      "true AP, an upper bound) and a *noise oracle* that chooses between raw and the null re-encode (showing how much "
      "gain pure label noise produces). Differences are tested with a paired bootstrap over images.")
    H2("3.12 Demo application design")
    P("The demo is a local web application with two parts. A FastAPI [@fastapi] back end (app/server.py) exposes a "
      "JSON API around the project code: the enhancers, UIQM/UCIQE, the Lab 1–5 operations (app/vision.py) and the "
      "frozen detector, all running on the CPU. A hand-written front end (app/static: HTML, CSS and JavaScript, no "
      "framework) presents the results in a black-and-white design: a black navigation bar, an image picker, large "
      "headings, monospaced numbers, greyscale SVG charts, detection boxes drawn as white outlines, and ▲/▼ markers "
      "instead of colours for 'helps' and 'hurts'. Sliders recompute results live, and any image can be enlarged.")
    P("The user picks one of twelve held-out RUOD test images, chosen to include cases where an enhancer helps, where "
      "all enhancers hurt, and where differences are within noise, or uploads any underwater image. The image is "
      "decoded and resized exactly as in the label pipeline, and four pages present the analysis:")
    BUL(["**Enhancement:** the raw image and the five enhanced versions with their colour histograms.",
         "**Quality scores:** live UIQM (with its UICM, UISM and UIConM components) and UCIQE for every version, next to "
         "the true detection change from the study's labels and the benchmark's reliability table.",
         "**Lab 1–5 operations:** OpenCV operations from Labs 1–5 on the raw or any enhanced version (Chapter 5), and a "
         "table comparing low-level features of all versions.",
         "**Detection:** the trained YOLO11-s detector on every version, with boxes drawn by cv2.rectangle, and, for "
         "sample images with ground truth, the per-image AP, AP50, F1 and ΔAP of each enhancer."])
    TAB("Demo application API.", ["Endpoint", "Returns"], [
        ["GET /api/status", "available samples, detector / FUnIE-GAN status, benchmark table"],
        ["POST /api/upload", "an id for an uploaded image"],
        ["GET /api/enhance/{id}", "the six versions and their colour histograms"],
        ["GET /api/quality/{id}", "UIQM components and UCIQE of every version, study ΔAP"],
        ["GET /api/lab/{id}/{1–5}", "Lab 1–5 operation results for a chosen version and parameters"],
        ["GET /api/features/{id}", "low-level feature statistics of every version"],
        ["GET /api/detect/{id}", "detections of every version, per-image AP, AP50, F1 and ΔAP"],
    ], widths=[2.2, 4.3], size=9.5)
    P("Expensive results are cached per image, and the model weights, FUnIE-GAN and the samples are packed by a Kaggle "
      "notebook into a single asset folder.")
    H2("3.13 Implementation and reproducibility")
    P("All code is in a Python package (src/uwiqa) with scripts and seven thin Kaggle notebooks; each notebook pulls the "
      "latest code from GitHub and calls one-line scripts, and long jobs are sharded and resumable. 74 unit tests check "
      "the metrics, correlation statistics, splits, detection AP and Lab operations.")
    TAB("Kaggle notebooks.", ["Notebook", "Content", "Hardware / time"], [
        ["01", "UID2021 download, 16 NR-IQA metrics, benchmark, stress test, diagnostics", "T4 / ≈1.5 h"],
        ["02", "RUOD/DUO verification, near-duplicate search", "CPU / ≈20 min"],
        ["03", "Leakage-safe splits, YOLO11-s training and held-out evaluation", "2 × T4 / ≈2.3 h"],
        ["04", "Per-image utility labels for 11,539 images × 7 versions", "T4 / ≈50 min"],
        ["05", "Backbone features, NR-IQA of every version, predictor, policies", "T4 / ≈3 h"],
        ["06", "Policies tuned on dataset mAP, paper figures", "CPU / ≈30 min"],
        ["07", "Assets for the demo application", "CPU / ≈5 min"],
    ], widths=[0.9, 4.0, 1.6], size=9.5)

    # ----------------------------------------------------------------- chapter 4
    H1("CHAPTER 4: RESULTS & SNAPSHOTS")
    H2("4.1 Reference detector")
    TAB("Frozen YOLO11-s detector.", ["Evaluation set", "Images", "mAP50", "mAP50:95"], [
        ["det-val", "700", "0.811", "0.569"],
        ["RUOD pred-test (held out)", "1,400", "0.798", "0.556"],
        ["DUO-clean (cross-dataset, 4 classes)", "4,539", "0.639", "0.387"],
    ], widths=[3.2, 1.0, 1.1, 1.2])
    P("The held-out score matches the validation score, so the detector is not over-fitted to its split. The weakest "
      "RUOD classes are corals (0.56 mAP50) and jellyfish (0.59). On DUO the drop comes mainly from echinus recall (0.54 "
      "at 0.96 precision; DUO has about nine small sea urchins per image). Recomputed from the stored detections at the "
      "1280-pixel working resolution, the scores are 0.799/0.561 (RUOD) and 0.644/0.402 (DUO), so the labelling "
      "protocol does not change the detector's behaviour.")
    H2("4.2 Data-leakage findings")
    P("The duplicate search found 7,979 near-duplicate pairs: **3,252 link RUOD to DUO** and **1,609 cross the public "
      "RUOD train/validation/test splits**. 3,190 of the 7,782 DUO images share a duplicate group with RUOD; only the "
      "remaining 4,539 labelled images were used as the cross-dataset test. Figure 4.1 shows examples.")
    FIG(ASSETS / "snap_dups_0_20.png", "Snapshot: identical frames found across RUOD and DUO (left/right, Hamming "
        "distance 0) and borderline duplicates across RUOD splits (distance 20).", 4.6)
    FIG(ASSETS / "snap_dups_21_26.png", "Snapshot: pairs at distance 21–26 crossing RUOD splits; all are frames of the "
        "same scene, so the threshold was raised to 32.", 4.6)
    H2("4.3 NR-IQA metrics vs. human opinion")
    rows = [["URanker", "0.623", "[0.581, 0.665]", "0.640", "0.445", "0.578"],
            ["LIQE-mix", "0.484", "[0.434, 0.529]", "0.480", "0.334", "0.381"],
            ["LIQE", "0.482", "[0.431, 0.528]", "0.499", "0.330", "0.401"],
            ["MUSIQ", "0.403", "[0.349, 0.458]", "0.422", "0.276", "0.303"],
            ["ARNIQA", "0.388", "[0.332, 0.443]", "0.408", "0.268", "0.289"],
            ["UCIQE", "0.334", "[0.275, 0.388]", "0.369", "0.227", "0.351"],
            ["UIQM", "0.325", "[0.272, 0.379]", "0.323", "0.218", "0.293"],
            ["TOPIQ-NR", "0.316", "[0.257, 0.373]", "0.330", "0.212", "0.167"],
            ["UCIQE (legacy)", "0.312", "[0.251, 0.367]", "0.366", "0.213", "0.386"],
            ["CLIP-IQA+", "0.298", "[0.238, 0.358]", "0.347", "0.201", "0.068"],
            ["TReS", "0.290", "[0.230, 0.344]", "0.306", "0.194", "0.070"],
            ["QualiCLIP", "0.249", "[0.187, 0.313]", "0.276", "0.169", "0.022"],
            ["MANIQA", "0.214", "[0.152, 0.273]", "0.241", "0.142", "−0.090"],
            ["CLIP-IQA", "0.190", "[0.129, 0.254]", "0.204", "0.129", "0.021"],
            ["NIQE", "0.158", "[0.095, 0.215]", "0.177", "0.106", "0.074"],
            ["BRISQUE", "0.034", "[−0.023, 0.095]", "0.072", "0.022", "−0.056"]]
    TAB("NR-IQA metrics vs. MOS on UID2021 (960 images). Every metric is significantly below URanker (paired bootstrap, "
        "p < 0.005).", ["Metric", "SRCC", "95% CI", "PLCC", "KRCC", "Intra-scene SRCC"], rows,
        widths=[1.5, 0.75, 1.4, 0.75, 0.75, 1.35], size=9.5)
    FIG(ASSETS / "fig_o1_srcc.png", "SRCC with MOS (bars, 95% CI) and mean intra-scene SRCC (diamonds).", 4.8)
    BUL(["**URanker, the only metric trained on underwater enhancements, is clearly best** (SRCC 0.62, intra-scene "
         "0.58); the best in-air models (LIQE, MUSIQ, ARNIQA) reach 0.39–0.48.",
         "**Intra-scene SRCC separates the metrics far more sharply.** CLIP-IQA(+), TReS, QualiCLIP and MANIQA fall to "
         "between −0.09 and 0.07 when ranking the versions of one scene: their global correlation comes mostly from "
         "differences between scenes. TOPIQ-NR reaches only 0.32 (0.17 within scenes), far below the ≈0.8 reported on a "
         "four-enhancer study [@palazari2026reliable].",
         "**UIQM and UCIQE reach only 0.33**, below the 0.54/0.60 reported by the UID2021 authors [@hou2023uid2021]. A "
         "diagnostic confirmed that all 960 image–score pairs align and that per-method mean MOS is sensible (UWCNN "
         "1.56 < raw 1.88 < … < UWB-VCSE 6.78). The remaining explanations are the revision of the public MOS from 52 "
         "to 77 observers (the original release is no longer available) or the original MATLAB implementations."])
    H2("4.4 Over-enhancement stress test")
    TAB("Spearman correlation between score and distortion strength (mean over 60 raw images). Large positive values "
        "mean the metric can be gamed.", ["Metric", "Contrast", "Red shift", "Saturation", "Unsharp"], [
        ["UCIQE", "0.91", "−0.01", "0.76", "1.00"], ["URanker", "0.89", "−0.51", "−0.86", "1.00"],
        ["MUSIQ", "0.08", "−0.10", "−0.63", "0.68"], ["TOPIQ-NR", "−0.34", "−0.01", "−0.52", "0.60"],
        ["LIQE", "0.28", "−0.18", "−0.47", "−0.10"], ["UIQM", "−0.77", "−0.30", "−0.73", "−0.09"]],
        widths=[1.6, 1.2, 1.2, 1.2, 1.2])
    FIG(ASSETS / "fig_o2_sweeps.png", "Mean change of each metric relative to the unprocessed image (in units of the "
        "metric's standard deviation) as distortion strength grows.", 6.4)
    P("UCIQE rises almost monotonically with contrast, saturation and sharpening, the very over-processing it is used to "
      "judge. URanker, the best metric against MOS, is also inflated by contrast stretching and unsharp masking, though "
      "not by colour distortion; MUSIQ and TOPIQ reward over-sharpening. UIQM rises slightly under mild red "
      "over-compensation before falling, which a single correlation hides.")
    H2("4.5 Effect of enhancement on detection")
    FIG(ASSETS / "snap_enhancers_grid.png", "Snapshot from the labelling run: raw held-out RUOD images and the five "
        "enhanced versions.", 6.4)
    TAB("Mean per-image AP50:95 and change vs. raw. 'helps' / 'hurts' count ΔAP > 0 / < 0 before noise calibration.",
        ["Method", "RUOD ΔAP", "helps", "hurts", "DUO ΔAP", "helps", "hurts"], [
        ["raw (AP)", "0.673", "", "", "0.531", "", ""],
        ["Null (JPEG q95)", "−0.001", "21%", "22%", "−0.001", "24%", "29%"],
        ["CLAHE", "−0.022", "27%", "50%", "−0.046", "33%", "64%"],
        ["UDCP", "−0.041", "24%", "54%", "−0.078", "26%", "71%"],
        ["FUnIE-GAN", "−0.066", "21%", "61%", "−0.107", "23%", "75%"],
        ["Gray-world", "−0.075", "20%", "60%", "−0.094", "26%", "72%"],
        ["Fusion", "−0.083", "20%", "62%", "−0.199", "14%", "84%"]],
        widths=[1.5, 0.9, 0.7, 0.7, 0.9, 0.7, 0.7], size=9.5)
    FIG(ASSETS / "fig_o3_delta.png", "Distribution of per-image ΔAP on RUOD; the grey band is the label-noise floor "
        "(±τ) estimated from the null re-encode.", 4.8)
    P("Every enhancer lowers mean detection AP, more so on DUO, which is out of distribution for the detector; this "
      "agrees with earlier studies [@wang2024isuie, @saleem2025influence]. The null re-encode is the crucial control: "
      "it 'helps' 21–24% of images, about as often as the enhancers do, so the sign of ΔAP alone is mostly noise. After "
      "calibration (ΔAP > τ = 0.05), each enhancer helps about 11% of RUOD images (9–19% on DUO).")
    H2("4.6 Predicting detection utility from the raw image")
    TAB("Predicting the per-image utility change (mean over five enhancers). Nothing is tuned on DUO.",
        ["Predictor", "RUOD SRCC", "RUOD AUROC", "RUOD BA", "DUO SRCC", "DUO AUROC"], [
        ["MLP, CLIP", "0.319", "0.597", "0.562", "0.174", "0.597"],
        ["MLP, CLIP+DINOv2", "0.314", "0.598", "0.576", "0.172", "0.602"],
        ["MLP, ResNet-18", "0.311", "0.597", "0.550", "0.158", "0.597"],
        ["MLP, DINOv2", "0.308", "0.578", "0.545", "0.156", "0.587"],
        ["MLP, CLIP, no ranking loss", "0.319", "0.595", "0.563", "0.168", "0.596"],
        ["Logistic, CLIP", "–", "0.596", "0.567", "–", "0.593"],
        ["Ridge, CLIP", "0.315", "0.519", "0.506", "0.154", "0.483"],
        ["Quality gain, UCIQE", "0.133", "0.532", "0.532", "−0.015", "0.442"],
        ["Quality gain, UIQM", "0.026", "0.515", "0.518", "−0.139", "0.466"],
        ["Quality gain, TOPIQ", "−0.029", "0.491", "0.507", "0.003", "0.469"],
        ["Quality gain, LIQE", "−0.130", "0.486", "0.502", "−0.003", "0.446"],
        ["Quality gain, URanker", "−0.182", "0.494", "0.509", "−0.025", "0.484"],
        ["Raw quality −Q(x), LIQE", "−0.235", "0.513", "0.518", "−0.026", "0.598"],
        ["Raw quality −Q(x), TOPIQ", "−0.212", "0.528", "0.513", "−0.020", "0.593"],
        ["Raw quality −Q(x), URanker", "−0.223", "0.499", "0.508", "−0.118", "0.503"]],
        widths=[2.2, 0.85, 0.9, 0.8, 0.85, 0.9], size=9)
    P("Learned raw-only predictors rank the utility change far better than any NR-IQA signal (SRCC 0.31–0.32 on RUOD and "
      "0.16–0.17 on DUO, against at most 0.13). The quality *gain* of an enhancement, the quantity practitioners "
      "optimise, is uncorrelated or negatively correlated with its detection gain for every metric, including URanker. "
      "Raw-image quality points the other way from the usual intuition: for TOPIQ, LIQE and URanker, −Q(x) has SRCC ≈ "
      "−0.22, i.e. enhancement hurts detection *more* on images that look worse. The choice of backbone and the ranking "
      "loss barely matter, which suggests that the limit is the signal rather than model capacity.")
    H2("4.7 Selective enhancement")
    TAB("Dataset mAP50:95 of enhancement policies (margins and model chosen on RUOD pred-val; nothing tuned on DUO).",
        ["Policy", "RUOD mAP", "Δ", "DUO mAP", "Δ"], [
        ["Always raw", "0.5548", "–", "0.4020", "–"],
        ["Always CLAHE (best on RUOD)", "0.5299", "−0.025", "0.3393", "−0.063"],
        ["Always fusion (worst)", "0.4550", "−0.100", "0.1946", "−0.207"],
        ["NR-IQA choice, UCIQE", "0.5087", "−0.046", "0.3363", "−0.066"],
        ["NR-IQA choice, TOPIQ", "0.4655", "−0.089", "0.2279", "−0.174"],
        ["NR-IQA, margin tuned (AP)", "0.5541", "−0.001", "0.3904", "−0.012"],
        ["NR-IQA, margin tuned (mAP)", "0.5547", "−0.000", "0.4014", "−0.001"],
        ["Learned, tuned (AP)", "0.5551", "+0.000", "0.4020", "0.000"],
        ["Learned, tuned (mAP)", "0.5531", "−0.002", "0.4009", "−0.001"],
        ["Noise oracle (control)", "0.5572", "+0.002", "0.4036", "+0.002"],
        ["Oracle (upper bound)", "0.5801", "+0.025", "0.4233", "+0.021"]],
        widths=[2.6, 1.0, 0.8, 1.0, 0.8], size=9.5)
    FIG(ASSETS / "fig_o5_policies_map.png", "Change in dataset mAP50:95 relative to always using the raw image "
        "(policies tuned on dataset mAP).", 6.4)
    P("The per-image oracle gains 0.025 (RUOD) and 0.021 (DUO) mAP over always-raw, about ten times the noise oracle, so "
      "real headroom exists, but no deployable policy recovers it. Tuned for mean per-image AP, the learned policy "
      "enhances 3% of images and ties always-raw (+0.0003, 95% CI [−0.0002, 0.0007]); tuned directly on dataset mAP it "
      "enhances 3.9% and loses a small but significant 0.0017 (CI [−0.0037, −0.0004]) on RUOD and 0.0010 on DUO. All "
      "tuned policies beat every always-enhance policy, but only because they learn to almost never enhance, whereas "
      "choosing the enhancer by NR-IQA gain costs 0.046–0.089 mAP on RUOD and 0.066–0.174 on DUO.")
    H2("4.8 Demo application snapshots")
    P("The following snapshots show the user interface of the demo application (Section 3.12) on held-out RUOD "
      "test images from the bundled samples.")
    for name, cap in [
            ("ui_overview", "Overview page with the navigation bar, image picker and key results"),
            ("ui_enhance", "Enhancement page: raw image, five enhanced versions and colour histograms"),
            ("ui_quality", "Quality page: live UIQM components and UCIQE next to the true detection change, "
                           "and the benchmark reliability table"),
            ("ui_labS", "Lab page, summary tab: low-level features of every version"),
            ("ui_detect", "Detection page: YOLO11-s detections on every version with per-image AP and ΔAP"),
            ("ui_detect_2", "Detection page for a second sample image")]:
        B.append(("shot", name, f"Application snapshot: {cap}"))
    P("Screenshots of the Lab 1–5 pages appear with the corresponding sections of Chapter 5.")
    H2("4.9 Summary of findings")
    BUL(["Only the underwater-trained URanker tracks human opinion reasonably (SRCC 0.62); in-air deep metrics cannot rank "
         "enhancements of one scene, and UCIQE can be gamed by over-processing.",
         "Every enhancer lowers mean detection AP, but each helps about one image in nine by more than the noise level.",
         "A raw-image predictor ranks detection utility better than any NR-IQA metric, but not well enough to turn the "
         "per-image headroom into a dataset-level gain; choosing enhancement by NR-IQA gain is harmful.",
         "Public RUOD/DUO splits leak near-duplicate frames; leakage-safe splits are needed for honest evaluation."])

    # ----------------------------------------------------------------- chapter 5
    H1("CHAPTER 5: LAB-1 TO LAB-5 FUNCTIONS UTILIZED")
    P("This chapter maps the topics of Labs 1–5 of the FCV Lab manual to where they are used. *Pipeline* means the "
      "research pipeline that produced the results of Chapter 4; *demo* means the application of Section 3.12. Theory "
      "follows the course references [@gonzalez2012, @szeliski2011]. Code excerpts are taken directly from the project "
      "source.")
    TAB("Lab 1–5 topics and where they are used.", ["Lab", "Topic (lab manual)", "Function(s)", "Where used"], [
        ["1", "Read / write images", "Image.open, cv2.imencode / imdecode", "Pipeline: data loading, null JPEG re-encode"],
        ["1", "Colour conversion", "cv2.cvtColor (RGB↔BGR, Lab, HSV, YCrCb, gray)", "Pipeline: CLAHE, UCIQE, sweeps; demo"],
        ["1", "RGB pixel values", "NumPy channel indexing", "Pipeline: UICM, red-shift sweep; demo pixel inspector"],
        ["1", "Resizing", "cv2.resize (INTER_AREA)", "Pipeline: 1280-px working size, 224-px backbone input"],
        ["1", "Rotating", "cv2.getRotationMatrix2D, cv2.warpAffine", "Demo"],
        ["1", "Drawing rectangles", "cv2.rectangle, cv2.putText", "Demo: detection boxes"],
        ["2", "Histogram", "cv2.calcHist", "Demo: colour and grey-level histograms"],
        ["2", "Histogram equalisation", "cv2.equalizeHist; cv2.createCLAHE", "Demo; pipeline enhancer (CLAHE)"],
        ["2", "Power-law (gamma)", "np.power", "Pipeline: fusion gamma input; demo"],
        ["2", "Piecewise-linear / contrast stretching", "linear stretch, min–max normalisation", "Pipeline: contrast sweep, fusion; demo"],
        ["2", "Contrast (Imax−Imin)/(Imax+Imin)", "block-wise contrast", "Pipeline: UIConM in UIQM; demo"],
        ["2", "Colour balance", "gray-world, red compensation", "Pipeline enhancers"],
        ["2", "Negative, log transform", "255 − r, c·log(1 + r)", "Demo"],
        ["3", "Gaussian blur", "cv2.GaussianBlur", "Pipeline: fusion, unsharp sweep; demo"],
        ["3", "Box vs. Gaussian filter", "cv2.boxFilter / cv2.blur", "Pipeline: guided filter in UDCP; demo"],
        ["3", "Median filter", "cv2.medianBlur", "Demo"],
        ["3", "Sobel gradient", "ndimage.sobel / cv2.Sobel", "Pipeline: UISM sharpness; demo"],
        ["3", "Laplacian", "cv2.Laplacian", "Pipeline: fusion weights; demo"],
        ["3", "Unsharp masking", "cv2.addWeighted", "Pipeline: stress sweep, fusion; demo"],
        ["3", "Canny edge detection", "cv2.Canny", "Demo: edge density raw vs. enhanced"],
        ["4", "Thresholding", "binary comparisons; cv2.threshold (Otsu), adaptiveThreshold",
         "Pipeline: dHash bits, dark-channel selection; demo"],
        ["4", "Region / block partition", "10×10 block statistics", "Pipeline: EME and logAMEE in UIQM"],
        ["4", "Region merging (grouping)", "union-find of duplicate pairs", "Pipeline: leakage-safe splits"],
        ["4", "K-means clustering", "cv2.kmeans", "Demo: colour segmentation"],
        ["5", "Feature extraction for representation learning", "frozen CLIP / DINOv2 / ResNet-18",
         "Pipeline: predictor input"],
        ["5", "Task-specific features", "colourfulness, sharpness, contrast", "Pipeline: UIQM / UCIQE"],
        ["5", "Binary (LBP-style) descriptors, matching", "dHash + Hamming distance", "Pipeline: near-duplicate retrieval"],
        ["5", "SIFT, FAST, Harris, HOG, LBP", "cv2.SIFT_create, FastFeatureDetector, cornerHarris, skimage hog / lbp",
         "Demo: features raw vs. enhanced"],
        ["5", "Precision, recall, F-measure", "per-image AP, F1", "Pipeline: detection utility"],
    ], widths=[0.4, 1.9, 2.2, 2.0], size=8.5)

    H2("5.1 Lab 1: Introduction to OpenCV")
    P("Images are read with PIL (with a fast reduced-size JPEG decode), converted between colour spaces with "
      "cv2.cvtColor and resized with cv2.resize using area interpolation. Every image is brought to a long side of 1280 "
      "pixels before enhancement and detection, so that all versions of an image are compared at the same "
      "resolution. The null enhancer writes and reads the image through an in-memory JPEG (cv2.imencode / imdecode). "
      "The demo adds pixel inspection, rotation and drawing of detection boxes.")
    CODE("Resizing (src/uwiqa/enhance/registry.py) and the JPEG re-encode used as the label-noise control.",
         src(registry.resize_long_side) + "\n\n" + src(registry._jpeg_roundtrip))
    CODE("Rotation and drawing rectangles in the demo (app/vision.py).",
         src(V.rotate) + "\n\n" + code_lines(V.draw_boxes, "out = img.copy()", "cv2.rectangle(out, (x1, y1)"))
    B.append(("shot", "ui_lab1", "Application snapshot: demo, Lab 1 tab (image information, pixel inspector, grayscale, rotation, resizing)"))
    H2("5.2 Lab 2: Image enhancement")
    P("Two of the five enhancers are Lab 2 techniques. CLAHE is a local, contrast-limited form of histogram "
      "equalisation applied to the lightness channel; gray-world is a colour-balance correction. The fusion enhancer "
      "uses a power-law (gamma) transformation and min–max contrast stretching, and the contrast measure of Lab 2, "
      "(Imax − Imin)/(Imax + Imin), is exactly the block contrast inside UIQM's UIConM term. The over-enhancement "
      "stress test uses linear contrast stretching. The demo shows histograms, global equalisation, negative, log, "
      "gamma and piecewise-linear transformations.")
    CODE("CLAHE and gray-world enhancers (src/uwiqa/enhance/classical.py).",
         src(classical.clahe) + "\n\n" + src(classical._gray_world))
    CODE("Michelson-type block contrast in UIQM (src/uwiqa/metrics/uiqm.py).", src(uiqm_mod.uiconm))
    CODE("Histogram equalisation and gamma in the demo (app/vision.py).", src(V.hist_equalize) + "\n\n" + src(V.gamma))
    B.append(("shot", "ui_lab2", "Application snapshot: demo, Lab 2 tab (negative, log, gamma, contrast stretching, histogram equalisation, CLAHE)"))
    H2("5.3 Lab 3: Image filtering")
    P("UIQM measures sharpness from the Sobel gradient magnitude of each colour channel. The fusion enhancer uses "
      "Gaussian blurring (for its sharpened input and saliency map) and the Laplacian (for its contrast weight), and "
      "blends its inputs with Gaussian and Laplacian pyramids. UDCP refines its transmission map with a guided filter "
      "built from box filters, and its dark channel is a minimum (erosion) filter. Unsharp masking is one of the four "
      "over-enhancement distortions. The demo compares box, Gaussian and median filters and shows Sobel, Laplacian and "
      "Canny [@canny1986] edges with their edge density.")
    CODE("Sobel gradient magnitude in UIQM's sharpness term (src/uwiqa/metrics/uiqm.py).", src(uiqm_mod._sobel_mag))
    CODE("Unsharp masking distortion (src/uwiqa/enhance/sweeps.py) and box-filter guided filter in UDCP "
         "(src/uwiqa/enhance/classical.py).", src(sweeps.unsharp) + "\n\n" + src(classical._guided_filter))
    B.append(("shot", "ui_lab3", "Application snapshot: demo, Lab 3 tab (box, Gaussian and median filters, Sobel, Laplacian, Canny, unsharp masking)"))
    H2("5.4 Lab 4: Image segmentation")
    P("Thresholding appears in three places: each dHash bit is a binary threshold of a pixel difference, UDCP selects "
      "the brightest 0.1% of its dark channel to estimate the background light, and the noise threshold τ turns the "
      "continuous ΔAP into a 'helps' label. UIQM partitions the image into 10×10 blocks and computes statistics per "
      "region, and the near-duplicate pairs are merged into connected groups with union-find, the graph analogue of "
      "region merging. The demo adds Otsu [@otsu1979] and adaptive thresholding and k-means colour segmentation.")
    CODE("Difference hash: binary comparisons of neighbouring pixels (src/uwiqa/data/dedup.py).", src(dedup.dhash))
    CODE("Otsu thresholding and k-means colour segmentation in the demo (app/vision.py).",
         src(V.otsu) + "\n\n" + src(V.kmeans_segment))
    B.append(("shot", "ui_lab4", "Application snapshot: demo, Lab 4 tab (Otsu and adaptive thresholding, k-means segmentation)"))
    H2("5.5 Lab 5: Feature extraction")
    P("Feature extraction is central to the predictor: every raw image is represented by a frozen deep feature vector "
      "(CLIP, DINOv2 or ResNet-18), the representation-learning view of Lab 5. UIQM and UCIQE are task-specific "
      "hand-crafted features (colourfulness, sharpness, contrast). The dHash descriptor is built like LBP from "
      "comparisons of neighbouring pixels and matched by Hamming distance for near-duplicate retrieval, and a unit test "
      "checks its robustness to resizing and re-compression. Detection quality is measured with precision, recall, F1 "
      "and AP. The demo computes SIFT [@lowe2004sift], FAST [@rosten2006fast], Harris [@harris1988], HOG "
      "[@dalal2005hog] and LBP [@ojala2002lbp] on raw and enhanced images and compares their counts and statistics.")
    CODE("Low-level feature summary used to compare raw and enhanced images in the demo (app/vision.py).",
         src(V.feature_summary))
    CODE("Per-image AP from precision and recall (src/uwiqa/detect/metrics.py).", src(dm._ap_from_tp))
    B.append(("shot", "ui_lab5", "Application snapshot: demo, Lab 5 tab (SIFT, FAST, Harris, HOG, LBP)"))
    H2("5.6 Lab topics not used")
    P("Video reading and writing (Lab 1), histogram specification (Lab 2), the Hough transform, region growing and region "
      "splitting (Lab 4) were not needed for this project; the image datasets are still images, and object localisation "
      "is performed by the trained detector.")

    # ----------------------------------------------------------------- chapter 6
    H1("CHAPTER 6: CONCLUSION, LIMITATIONS & FUTURE WORK")
    H2("6.1 Conclusion")
    P("Whether an enhanced underwater image looks better, to a human or to a metric, says little about whether a "
      "detector sees it better. On UID2021 only an underwater-trained ranker tracks human opinion reasonably well; in-air "
      "deep metrics cannot rank enhancements of one scene, and the de facto metric UCIQE can be gamed. For detection, "
      "every enhancer tested hurts on average, and selecting enhancement by NR-IQA gain hurts more, while a learned "
      "raw-image predictor at least avoids the harm. Per-image headroom exists, about ten times the noise level, but "
      "neither NR-IQA metrics nor frozen-feature predictors capture it yet. The leakage-safe splits, noise-calibrated "
      "per-image utility labels, code and demo application make these effects measurable rather than assumed.")
    H2("6.2 Limitations")
    BUL(["**Detector dependence:** the labels come from a detector trained on raw images, for which enhanced inputs are "
         "slightly out of distribution; this may overstate the harm of enhancement.",
         "**Enhancers:** four of the five enhancers are compact re-implementations, and the fusion port is untuned; "
         "TensorFlow-1 methods (Water-Net, Ucolor) were excluded.",
         "**Human data:** the benchmark uses one MOS dataset, whose released scores differ from those in its paper; the "
         "stress test shows score inflation but not human preference at each strength.",
         "**Per-image AP** is noisy for images with few objects; it is controlled with the null enhancer rather than "
         "removed.",
         "**Demo:** uploaded images have no ground truth, so AP and ΔAP can only be shown for the bundled samples."])
    H2("6.3 Future work")
    BUL(["Retrain the detector with enhancement augmentation (protocol (b)) and repeat the labelling.",
         "Extend the benchmark to SAUD and other subjective datasets, and run a small paired-comparison study on the "
         "stress-test images.",
         "Fine-tune the backbone, or predict utility per region or per class instead of per image.",
         "Add more learned enhancers and deploy the selective policy on an underwater vehicle pipeline."])

    # ----------------------------------------------------------------- chapter 7
    H1("CHAPTER 7: REFERENCES")
    B.append(("refs",))
    return B


# ===================================================================== docx helpers
def set_run_font(run, size=None, bold=None, italic=None, name=FONT, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def style(doc, name, size, bold=False, italic=False, align=None, before=0, after=6, line=1.5, outline=None,
          page_break=False, keep_next=False, font=FONT):
    st = doc.styles[name] if name in [s.name for s in doc.styles] else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = None
    f = st.font
    f.name, f.size, f.bold, f.italic = font, Pt(size), bold, italic
    f.color.rgb = RGBColor(0, 0, 0)
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), font)
    pf = st.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line
    if align is not None:
        pf.alignment = align
    pf.page_break_before = page_break
    pf.keep_with_next = keep_next
    ppr = st.element.get_or_add_pPr()
    for old in ppr.findall(qn("w:outlineLvl")):
        ppr.remove(old)
    if outline is not None:
        o = OxmlElement("w:outlineLvl")
        o.set(qn("w:val"), str(outline))
        ppr.append(o)
    return st


TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+`)")


def add_rich(par, text, size=None, base_italic=False):
    for tok in TOKEN.split(text):
        if not tok:
            continue
        if tok.startswith("**"):
            r = par.add_run(tok[2:-2])
            set_run_font(r, size=size, bold=True, italic=base_italic or None)
        elif tok.startswith("`"):
            r = par.add_run(tok[1:-1])
            set_run_font(r, size=(size or 12) - 1, name="Consolas")
        elif tok.startswith("*") and len(tok) > 2:
            r = par.add_run(tok[1:-1])
            set_run_font(r, size=size, italic=True)
        else:
            r = par.add_run(tok)
            set_run_font(r, size=size, italic=base_italic or None)


def shade(element, fill):
    pr = element.get_or_add_tcPr() if element.tag.endswith("tc") else element.get_or_add_pPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), fill)
    pr.append(sh)


def table_borders(table, color="8C8C8C", size=4):
    tblpr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(size))
        e.set(qn("w:color"), color)
        b.append(e)
    tblpr.append(b)


def add_field(par, instr, placeholder=""):
    r = par.add_run()
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    r._r.append(b)
    r2 = par.add_run()
    t = OxmlElement("w:instrText")
    t.set(qn("xml:space"), "preserve")
    t.text = instr
    r2._r.append(t)
    r3 = par.add_run()
    s = OxmlElement("w:fldChar")
    s.set(qn("w:fldCharType"), "separate")
    r3._r.append(s)
    r4 = par.add_run(placeholder)
    r5 = par.add_run()
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    r5._r.append(e)
    return [r, r2, r3, r4, r5]


def delete_paragraph(p):
    p._element.getparent().remove(p._element)


def replace_runs(par, texts: list[tuple[str, dict]]):
    """Replace all runs of a paragraph by the given (text, font kwargs) runs, keeping paragraph formatting."""
    size = next((r.font.size.pt for r in par.runs if r.font.size), 14)
    for r in list(par.runs):
        r._element.getparent().remove(r._element)
    for t, kw in texts:
        run = par.add_run(t)
        set_run_font(run, size=kw.get("size", size), bold=kw.get("bold"), italic=kw.get("italic"))
        if kw.get("break"):
            run.add_break()


# ===================================================================== assets
def make_pipeline_figure(path: Path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch

    fig, ax = plt.subplots(figsize=(10, 4.6))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 4.6)
    ax.axis("off")

    def box(x, y, w, h, text, fc="#eef3fb"):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.02,rounding_size=0.08",
                                    fc=fc, ec="#2a78d6", lw=1.2))
        ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=9, family="serif")

    def arrow(x1, y1, x2, y2, style="-|>", ls="-"):
        ax.annotate("", xy=(x2, y2), xytext=(x1, y1),
                    arrowprops=dict(arrowstyle=style, color="#52514e", lw=1.2, linestyle=ls))

    ax.text(0.05, 4.35, "Benchmark (UID2021)", fontsize=9.5, weight="bold", family="serif")
    box(0.1, 3.45, 1.7, 0.7, "960 enhanced\nimages + MOS")
    box(2.3, 3.45, 2.1, 0.7, "16 NR-IQA metrics\n(UIQM, UCIQE, deep)")
    box(4.9, 3.45, 2.3, 0.7, "SRCC / PLCC / KRCC,\nbootstrap, intra-scene")
    box(7.7, 3.45, 2.2, 0.7, "stress test\n(over-enhancement)")
    for a, b in ((1.8, 2.3), (4.4, 4.9), (7.2, 7.7)):
        arrow(a, 3.8, b, 3.8)

    ax.text(0.05, 2.75, "Detection-utility labels (RUOD / DUO)", fontsize=9.5, weight="bold", family="serif")
    box(0.1, 1.85, 1.7, 0.7, "raw image x\n(leakage-safe split)")
    box(2.3, 1.85, 2.1, 0.7, "enhancer m\n(5 + null re-encode)")
    box(4.9, 1.85, 2.3, 0.7, "frozen YOLO11-s\n(trained on raw)")
    box(7.7, 1.85, 2.2, 0.7, "per-image AP,\nΔAPm(x)")
    for a, b in ((1.8, 2.3), (4.4, 4.9), (7.2, 7.7)):
        arrow(a, 2.2, b, 2.2)

    ax.text(0.05, 1.15, "Deployable decision", fontsize=9.5, weight="bold", family="serif")
    box(2.3, 0.2, 2.1, 0.7, "frozen backbone\n(CLIP / DINOv2 / ResNet)")
    box(4.9, 0.2, 2.3, 0.7, "head: predicted ΔAPm,\nP(helps m)")
    box(7.7, 0.2, 2.2, 0.7, "policy: apply m*\nor keep raw", fc="#e9f7f1")
    arrow(0.95, 1.85, 0.95, 0.55)
    arrow(0.95, 0.55, 2.3, 0.55)
    arrow(4.4, 0.55, 4.9, 0.55)
    arrow(7.2, 0.55, 7.7, 0.55)
    arrow(8.8, 1.85, 8.8, 0.9, ls="--")
    ax.text(8.9, 1.38, "training\nlabels", fontsize=8, family="serif", color="#52514e")
    fig.savefig(path, dpi=220, bbox_inches="tight")
    plt.close(fig)


def split_tall(path: Path, max_ratio: float = 1.3) -> list[Path]:
    """Cut a tall screenshot at blank rows into equal-ish parts of height <= max_ratio * width."""
    im = np.asarray(Image.open(path).convert("RGB"))
    rows = np.where(im.std(axis=(1, 2)) >= 1.0)[0]
    im = im[: rows[-1] + 24] if len(rows) else im  # trim the blank bottom margin
    h, w = im.shape[:2]
    n = int(np.ceil(h / (max_ratio * w)))
    blank = np.where(im.std(axis=(1, 2)) < 1.0)[0]
    cuts = [0]
    for k in range(1, n):
        target = k * h // n
        near = blank[(blank > cuts[-1]) & (np.abs(blank - target) < 0.2 * w)]
        cuts.append(int(near[np.argmin(np.abs(near - target))]) if len(near) else target)
    cuts.append(h)
    out = []
    for k in range(n):
        p = path.with_name(f"{path.stem}_part{k + 1}.png")
        Image.fromarray(im[cuts[k]:cuts[k + 1]]).save(p)
        out.append(p)
    return out


# ===================================================================== build
def build(out_docx: Path):
    blocks = content()
    # citation numbers in order of first appearance
    order = []
    cite_re = re.compile(r"\[(@[^\]]+)\]")

    def collect(text):
        for m in cite_re.finditer(text):
            for k in m.group(1).split(","):
                k = k.strip().lstrip("@")
                if k not in REFS:
                    raise KeyError(f"unknown reference {k}")
                if k not in order:
                    order.append(k)

    for b in blocks:
        for item in b[1:]:
            if isinstance(item, str):
                collect(item)
            elif isinstance(item, list):
                for x in item:
                    if isinstance(x, str):
                        collect(x)
                    elif isinstance(x, list):
                        for y in x:
                            collect(str(y))
    num = {k: i + 1 for i, k in enumerate(order)}

    def cite(text):
        return cite_re.sub(lambda m: "[" + ", ".join(str(num[k.strip().lstrip("@")])
                                                     for k in m.group(1).split(",")) + "]", text)

    doc = Document(HERE / "template" / "FCV_MiniProject_Template.docx")
    P = doc.paragraphs

    # ---- cover page
    replace_runs(P[9], [(TITLE, {"size": 20, "bold": True})])
    P[9].alignment = WD_ALIGN_PARAGRAPH.CENTER
    lines = [(f"{n}  (Reg. No. {r}, Roll No. {roll})", {"size": 13, "break": True}) for n, r, roll in TEAM]
    lines.append((f"Section: {SECTION}", {"size": 13}))
    replace_runs(P[18], lines)
    P[18].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i in (10, 11, 12):  # fewer blank lines under the longer title
        if not P[i].text.strip():
            delete_paragraph(P[i])
    P = doc.paragraphs
    names_i = next(i for i, p in enumerate(P) if p.text.startswith(TEAM[0][0]))
    school_i = next(i for i, p in enumerate(P) if p.text.strip() == "School of Computer Engineering")
    blanks = [p for p in P[names_i + 1:school_i] if not p.text.strip()]
    for p in blanks[2:]:  # keep two spacer lines so the cover stays on one page
        delete_paragraph(p)
    P = doc.paragraphs

    # ---- certificate
    cert = next(p for p in P if p.text.startswith("This is to certify"))
    names = ", ".join(f"{n} (Reg. No. {r})" for n, r, _ in TEAM[:-1]) + f" and {TEAM[-1][0]} (Reg. No. {TEAM[-1][1]})"
    replace_runs(cert, [
        ("This is to certify that the project titled ", {"bold": False}),
        (TITLE, {"bold": True}),
        (" is a record of the bonafide work done by ", {"bold": False}),
        (names, {"bold": True}),
        (" submitted in partial fulfilment of the requirements for the award of the Degree of Bachelor of Technology "
         "(B.Tech.) in Computer Science and Engineering (Artificial Intelligence and Machine Learning) of Manipal "
         "Institute of Technology, Manipal, Karnataka, (A Constituent Institute of Manipal Academy of Higher Education), "
         "during the academic year 2026-2027.", {"bold": False}),
    ])
    date_p = next(p for p in P if p.text.strip() == "00/00/2026")
    replace_runs(date_p, [("__/__/2026", {})])

    # ---- contents: replace the static list with a real TOC field
    toc_head = next(i for i, p in enumerate(P) if p.text.strip() == "TABLE OF CONTENTS")
    for p in P[toc_head + 1:]:
        delete_paragraph(p)

    style(doc, "Body FCV", 12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6, line=1.5)
    style(doc, "Heading 1", 16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=18, line=1.15,
          outline=0, page_break=True, keep_next=True)
    style(doc, "Heading 2", 13, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=12, after=6, line=1.15,
          outline=1, keep_next=True)
    style(doc, "Caption FCV", 10, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=12, line=1.0)
    style(doc, "Code FCV", 8, align=WD_ALIGN_PARAGRAPH.LEFT, before=2, after=2, line=1.0, font="Consolas")
    style(doc, "List FCV", 12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=3, line=1.4)

    tp = doc.add_paragraph()
    add_field(tp, 'TOC \\o "1-2" \\h \\z \\u', "Right-click and choose Update Field to build the table of contents.")

    # ---- body section with page numbers starting at 1
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.footer.is_linked_to_previous = False
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in add_field(fp, "PAGE", "1"):
        set_run_font(r, size=11)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "1")
    sec._sectPr.append(pg)
    doc.sections[0].footer.is_linked_to_previous = False

    chapter, fig_n, tab_n, first_h1 = 0, 0, 0, True
    for b in blocks:
        kind = b[0]
        if kind in ("h1", "h1u"):
            p = doc.add_paragraph(style="Heading 1")
            if first_h1:
                p.paragraph_format.page_break_before = False  # new section already starts a page
                first_h1 = False
            add_rich(p, b[1], size=16)
            if kind == "h1":
                chapter += 1
                fig_n = tab_n = 0
        elif kind == "h2":
            p = doc.add_paragraph(style="Heading 2")
            add_rich(p, b[1], size=13)
        elif kind == "p":
            p = doc.add_paragraph(style="Body FCV")
            add_rich(p, cite(b[1]), size=12)
        elif kind == "bullets":
            for item in b[1]:
                p = doc.add_paragraph(style="List FCV")
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.first_line_indent = Inches(-0.2)
                add_rich(p, "•  " + cite(item), size=12)
        elif kind == "table":
            _, cap, head, rows, widths, size = b
            tab_n += 1
            c = doc.add_paragraph(style="Caption FCV")
            c.paragraph_format.keep_with_next = True
            c.paragraph_format.space_after = Pt(4)
            add_rich(c, f"Table {chapter}.{tab_n}: {cite(cap)}", size=10, base_italic=True)
            t = doc.add_table(rows=1 + len(rows), cols=len(head))
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            table_borders(t)
            for j, h in enumerate(head):
                cell = t.rows[0].cells[j]
                cell.paragraphs[0].text = ""
                add_rich(cell.paragraphs[0], f"**{h}**", size=size)
                shade(cell._tc, "E6E6E6")
            for i, row in enumerate(rows):
                for j, v in enumerate(row):
                    cell = t.rows[i + 1].cells[j]
                    add_rich(cell.paragraphs[0], cite(str(v)), size=size)
            for row in t.rows:
                for j, cell in enumerate(row.cells):
                    pf = cell.paragraphs[0].paragraph_format
                    pf.space_after, pf.space_before, pf.line_spacing = Pt(1), Pt(1), 1.0
                    if widths:
                        cell.width = Inches(widths[j])
            doc.add_paragraph(style="Body FCV").paragraph_format.space_after = Pt(2)
        elif kind == "fig":
            _, path, cap, w = b
            if not Path(path).exists():
                print(f"  ! missing figure {path}")
                continue
            fig_n += 1
            p = doc.add_paragraph(style="Caption FCV")
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(path), width=Inches(w))
            c = doc.add_paragraph(style="Caption FCV")
            add_rich(c, f"Figure {chapter}.{fig_n}: {cite(cap)}", size=10, base_italic=True)
        elif kind == "shot":
            _, stem, cap = b
            tall = ASSETS / f"{stem}.png"
            if not tall.exists():
                print(f"  ! missing figure {tall}")
                continue
            parts = split_tall(tall)
            for k, part in enumerate(parts):
                fig_n += 1
                p = doc.add_paragraph(style="Caption FCV")
                p.paragraph_format.keep_with_next = True
                p.add_run().add_picture(str(part), width=Inches(6.0))
                c = doc.add_paragraph(style="Caption FCV")
                part_txt = f" (part {k + 1} of {len(parts)})" if len(parts) > 1 else ""
                add_rich(c, f"Figure {chapter}.{fig_n}: {cap}{part_txt}.", size=10, base_italic=True)
        elif kind == "code":
            _, cap, code = b
            fig_n += 1
            p = doc.add_paragraph(style="Code FCV")
            p.paragraph_format.left_indent = Inches(0.1)
            shade(p._p, "F4F4F2")
            lines = code.splitlines()
            for k, line in enumerate(lines):
                r = p.add_run(line)
                set_run_font(r, size=8, name="Consolas")
                if k < len(lines) - 1:
                    r.add_break()
            c = doc.add_paragraph(style="Caption FCV")
            add_rich(c, f"Listing {chapter}.{fig_n}: {cite(cap)}", size=10, base_italic=True)
        elif kind == "refs":
            for k in order:
                p = doc.add_paragraph(style="Body FCV")
                p.paragraph_format.left_indent = Inches(0.45)
                p.paragraph_format.first_line_indent = Inches(-0.45)
                p.paragraph_format.line_spacing = 1.15
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_rich(p, f"[{num[k]}]\t" + REFS[k], size=11)
    doc.save(out_docx)
    print(f"wrote {out_docx} ({len(order)} references)")


def word_finalize(docx_path: Path):
    """Update the TOC and page fields in Word and export a PDF (Windows + Word only)."""
    try:
        import win32com.client  # noqa: F401
    except ImportError:
        print("  pywin32 not available: open the .docx in Word and update the table of contents (F9).")
        return
    import win32com.client
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        d = word.Documents.Open(str(docx_path))
        for name in ("TOC 1", "TOC 2"):
            try:
                st = d.Styles(name)
                st.Font.Name = FONT
                st.Font.Size = 12 if name == "TOC 1" else 11
                st.Font.Bold = name == "TOC 1"
            except Exception:
                pass
        d.TablesOfContents(1).Update()
        d.Fields.Update()
        d.Save()
        pdf = docx_path.with_suffix(".pdf")
        d.ExportAsFixedFormat(str(pdf), 17)
        print(f"wrote {pdf} ({d.ComputeStatistics(2)} pages)")
        d.Close(0)
    finally:
        word.Quit()


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(HERE / "FCV_MiniProject_Report.docx"))
    out = Path(ap.parse_args().out)
    try:
        out.open("ab").close()
    except PermissionError:
        out = out.with_name(out.stem + "_new.docx")
        print(f"  ! report is open in another program; writing {out.name} instead")
    make_pipeline_figure(ASSETS / "pipeline.png")
    build(out)
    word_finalize(out)
